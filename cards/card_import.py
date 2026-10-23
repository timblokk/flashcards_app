import csv
from io import BytesIO, StringIO
from zipfile import BadZipFile

from django.core.exceptions import ValidationError


MAX_IMPORT_FILE_SIZE = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000


def _text_cell(value, row_number, column_name):
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ValidationError(
            f"Row {row_number}: {column_name} must contain text."
        )
    return str(value).strip()


def _validated_rows(rows):
    try:
        header = next(rows)
    except StopIteration:
        raise ValidationError("The file is empty.")

    normalized_header = [
        str(value).strip().casefold() if value is not None else ""
        for value in header
    ]
    if normalized_header not in (
        ["question", "answer"],
        ["question", "answer", "language"],
    ):
        raise ValidationError(
            "The first row must contain question, answer, and optionally language."
        )
    has_language = len(normalized_header) == 3

    parsed_rows = []
    for row_number, row in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        expected_columns = 3 if has_language else 2
        if len(row) != expected_columns:
            raise ValidationError(
                f"Row {row_number}: expected exactly {expected_columns} columns."
            )

        question = _text_cell(row[0], row_number, "question")
        answer = _text_cell(row[1], row_number, "answer")
        language = _text_cell(row[2], row_number, "language") if has_language else None
        if not question or not answer:
            raise ValidationError(
                f"Row {row_number}: question and answer must both be non-empty."
            )
        if has_language and not language:
            raise ValidationError(f"Row {row_number}: language must be non-empty.")
        parsed_rows.append((row_number, question, answer, language))
        if len(parsed_rows) > MAX_IMPORT_ROWS:
            raise ValidationError(
                f"Files may contain at most {MAX_IMPORT_ROWS} non-empty data rows."
            )

    if not parsed_rows:
        raise ValidationError("The file contains no card rows.")
    return parsed_rows


def parse_card_import(uploaded_file):
    if uploaded_file.size > MAX_IMPORT_FILE_SIZE:
        raise ValidationError("Files must be 5 MB or smaller.")

    filename = uploaded_file.name.casefold()
    content = uploaded_file.read()
    if filename.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValidationError("CSV files must use UTF-8 encoding.") from error
        try:
            return _validated_rows(csv.reader(StringIO(text, newline="")))
        except csv.Error as error:
            raise ValidationError(f"Could not read CSV file: {error}") from error

    if filename.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except (BadZipFile, OSError, ValueError, ImportError) as error:
            raise ValidationError("Could not read this .xlsx file.") from error
        try:
            worksheet = workbook.worksheets[0]
            return _validated_rows(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()

    raise ValidationError("Upload a .csv or .xlsx file.")
