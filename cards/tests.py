import json
from html.parser import HTMLParser
from io import BytesIO
from datetime import timedelta
from io import StringIO
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.contrib.auth import get_user_model
from django.db import IntegrityError, connection, transaction
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.migrations.executor import MigrationExecutor
from django.test import TestCase, TransactionTestCase
from django.template.loader import render_to_string
from django.utils import timezone
from django.urls import NoReverseMatch, reverse

from cards.forms import CardForm
from cards.models import Card, DictionaryEntry


class ArabicDisplayParser(HTMLParser):
	def __init__(self):
		super().__init__()
		self.displays = []
		self._active_display = None
		self._depth = 0

	def handle_starttag(self, tag, attrs):
		attributes = dict(attrs)
		if self._active_display is not None:
			self._depth += 1
		elif "arabic-display" in attributes.get("class", "").split():
			self._active_display = {
				"id": attributes.get("id"),
				"text": [],
			}
			self._depth = 1

	def handle_data(self, data):
		if self._active_display is not None:
			self._active_display["text"].append(data)

	def handle_endtag(self, tag):
		if self._active_display is None:
			return
		self._depth -= 1
		if self._depth == 0:
			self.displays.append({
				"id": self._active_display["id"],
				"text": " ".join("".join(self._active_display["text"]).split()),
			})
			self._active_display = None


def arabic_display_texts(response, element_id):
	parser = ArabicDisplayParser()
	parser.feed(response.content.decode())
	return [
		display["text"]
		for display in parser.displays
		if display["id"] == element_id
	]


class DictionarySuggestionTests(TestCase):
	def setUp(self):
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Apple",
			translation="التفاحة",
			vocalized_translation="التُّفَّاحَةُ",
		)
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Application",
			translation="تطبيق",
			vocalized_translation="تَطْبِيقٌ",
		)
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Water",
			translation="الماء",
			vocalized_translation="الْمَاءُ",
		)

	def test_suggests_prefix_matches_with_translations(self):
		response = self.client.get(reverse("dictionary-suggest"), {"q": "app"})

		self.assertEqual(response.status_code, 200)
		self.assertEqual(
			response.json()["results"],
			[
				{"id": 1, "headword": "Apple", "translation": "التفاحة", "vocalized_translation": "التُّفَّاحَةُ"},
				{"id": 2, "headword": "Application", "translation": "تطبيق", "vocalized_translation": "تَطْبِيقٌ"},
			],
		)

	def test_does_not_search_for_one_character(self):
		response = self.client.get(reverse("dictionary-suggest"), {"q": "a"})

		self.assertEqual(response.json(), {"results": []})

	def test_dictionary_suggestion_fills_question_and_answer(self):
		self.client.force_login(get_user_model().objects.create_user("card-owner"))
		response = self.client.post(
			reverse("card-create"),
			{
				"question": "Apple",
				"answer": "التفاحة",
				"dictionary_entry_id": DictionaryEntry.objects.get(headword="Apple").id,
				"box": 1,
			},
		)

		self.assertRedirects(response, reverse("card-create"))
		card = Card.objects.get(question="Apple")
		self.assertEqual(card.answer, "التفاحة")
		self.assertEqual(card.answer_vocalized, "التُّفَّاحَةُ")
		self.assertEqual(card.user.username, "card-owner")

	def test_dictionary_suggestion_allows_custom_vocalized_answer(self):
		self.client.force_login(get_user_model().objects.create_user("custom-vocalization-owner"))
		entry = DictionaryEntry.objects.get(headword="Apple")
		response = self.client.post(
			reverse("card-create"),
			{
				"question": "Apple",
				"answer": "التفاحة",
				"answer_vocalized": "تَفَّاحَةٌ",
				"dictionary_entry_id": entry.id,
				"box": 1,
			},
		)

		self.assertRedirects(response, reverse("card-create"))
		card = Card.objects.get(question="Apple")
		self.assertEqual(card.answer_vocalized, "تَفَّاحَةٌ")


class AccountDeckTests(TestCase):
	def create_card(self, **kwargs):
		return Card.objects.create(
			answer="legacy-secret-answer",
			question="question",
			**kwargs,
		)

	def test_card_pages_require_login(self):
		response = self.client.get(reverse("card-list"))

		self.assertRedirects(
			response,
			f"{reverse('login')}?next={reverse('card-list')}",
		)

	def test_login_redirects_to_card_list(self):
		user = get_user_model().objects.create_user(
			"login-redirect-user",
			password="Correct-Horse-92!",
		)

		response = self.client.post(
			reverse("login"),
			{"username": user.username, "password": "Correct-Horse-92!"},
		)

		self.assertRedirects(response, reverse("card-list"))

	def test_new_account_creates_language_neutral_cards_without_category(self):
		self.client.post(
			reverse("signup"),
			{
				"username": "new-user",
				"password1": "Correct-Horse-92!",
				"password2": "Correct-Horse-92!",
			},
		)
		user = get_user_model().objects.get(username="new-user")

		create_response = self.client.get(reverse("card-create"))
		self.assertContains(create_response, 'name="question"')
		self.assertContains(create_response, 'name="answer"')
		self.assertContains(create_response, 'name="answer_vocalized"')
		self.assertNotContains(create_response, 'name="category"')
		self.assertNotContains(create_response, 'name="language"')
		self.assertNotContains(create_response, 'name="question_en"')

		response = self.client.post(
			reverse("card-create"),
			{
				"question": "Apple",
				"answer": "التفاحة",
				"box": 1,
			},
		)
		self.assertRedirects(response, reverse("card-create"))
		self.assertEqual(Card.objects.get(user=user).question, "Apple")

	def test_signup_shows_password_validation_errors(self):
		response = self.client.get(reverse("signup"))

		self.assertContains(response, "Password requirements")
		self.assertContains(response, "at least 8 characters")
		self.assertContains(response, "commonly used password")
		self.assertContains(response, "entirely numeric")
		self.assertContains(response, "too similar to your other personal information")

		response = self.client.post(
			reverse("signup"),
			{
				"username": "new-user",
				"password1": "short",
				"password2": "short",
			},
		)

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, "This password is too short")
		self.assertContains(response, 'role="alert"')
		self.assertFalse(
			get_user_model().objects.filter(username="new-user").exists()
		)

	def test_card_form_prevents_duplicate_question_per_user(self):
		user = get_user_model().objects.create_user("owner")
		other_user = get_user_model().objects.create_user("other")
		Card.objects.create(user=user, question="question", answer="translation")
		duplicate_form = CardForm(
			data={
				"question": "question",
				"answer": "different translation",
				"box": 1,
			},
			user=user,
		)
		self.assertFalse(duplicate_form.is_valid())
		self.assertIn("already exists", str(duplicate_form.errors))

		other_user_form = CardForm(
			data={
				"question": "question",
				"answer": "translation",
				"box": 1,
			},
			user=other_user,
		)
		self.assertTrue(other_user_form.is_valid())

	def test_database_prevents_duplicate_question_per_user(self):
		user = get_user_model().objects.create_user("unique-question-owner")
		Card.objects.create(user=user, question="question", answer="translation")
		with self.assertRaises(IntegrityError), transaction.atomic():
			Card.objects.create(
				user=user,
				question="question",
				answer="different translation",
			)

	def test_legacy_username_claims_existing_cards_at_signup(self):
		legacy_card = self.create_card()

		response = self.client.post(
			reverse("signup"),
			{
				"username": "legacyUser",
				"password1": "Correct-Horse-92!",
				"password2": "Correct-Horse-92!",
			},
		)

		self.assertRedirects(response, reverse("card-list"))
		user = get_user_model().objects.get(username="legacyUser")
		self.assertTrue(user.check_password("Correct-Horse-92!"))
		legacy_card.refresh_from_db()
		self.assertEqual(legacy_card.user, user)
		self.assertIn("_auth_user_id", self.client.session)

	def test_other_accounts_start_with_empty_private_decks(self):
		legacy_card = self.create_card()
		self.client.post(
			reverse("signup"),
			{
				"username": "legacyUser",
				"password1": "Correct-Horse-92!",
				"password2": "Correct-Horse-92!",
			},
		)
		legacy_user = get_user_model().objects.get(username="legacyUser")

		other_client = self.client_class()
		response = other_client.post(
			reverse("signup"),
			{
				"username": "anotherUser",
				"password1": "A-Different-Password-75!",
				"password2": "A-Different-Password-75!",
			},
		)

		self.assertEqual(response.status_code, 302)
		other_user = get_user_model().objects.get(username="anotherUser")
		self.assertEqual(
			Card.objects.filter(user=other_user).count(),
			0,
		)

		response = other_client.get(reverse("card-list"))
		self.assertNotContains(response, "legacy-secret-answer")
		self.assertNotContains(response, "<mark>1</mark>")

		self.client.force_login(legacy_user)
		response = self.client.get(reverse("card-list"))
		self.assertContains(response, "legacy-secret-answer")
		legacy_card.refresh_from_db()
		self.assertEqual(legacy_card.user, legacy_user)

	def test_account_cannot_edit_or_delete_another_users_card(self):
		owner = get_user_model().objects.create_user("owner")
		other_user = get_user_model().objects.create_user("other")
		card = self.create_card(user=owner)
		self.client.force_login(other_user)

		self.assertEqual(
			self.client.get(reverse("card-update", args=[card.pk])).status_code,
			404,
		)
		self.assertEqual(
			self.client.post(reverse("card-delete", args=[card.pk])).status_code,
			404,
		)
		self.assertEqual(
			self.client.post(
				reverse("todo-list"),
				{"card_id": card.pk, "solved": "on"},
			).status_code,
			404,
		)
		self.assertTrue(Card.objects.filter(pk=card.pk).exists())

	def test_box_and_repeat_pages_only_show_the_current_users_cards(self):
		owner = get_user_model().objects.create_user("owner")
		other_user = get_user_model().objects.create_user("other")
		card = self.create_card(user=owner, box=1)
		card.date_moved = timezone.now() - timedelta(days=2)
		card.save(update_fields=["date_moved"])

		self.client.force_login(other_user)
		box_response = self.client.get(reverse("box", args=[1]))
		repeat_response = self.client.get(reverse("todo-list"))

		self.assertNotContains(box_response, "legacy-secret-answer")
		self.assertContains(box_response, "<mark>0</mark>")
		self.assertNotContains(repeat_response, "legacy-secret-answer")
		self.assertContains(repeat_response, "<mark>0</mark>")

		self.client.force_login(owner)
		self.assertContains(self.client.get(reverse("box", args=[1])), "legacy-secret-answer")
		self.assertContains(self.client.get(reverse("todo-list")), "legacy-secret-answer")

	def test_home_page_shows_all_user_cards_without_filter_view(self):
		user = get_user_model().objects.create_user("home-page-owner")
		self.create_card(user=user)
		self.client.force_login(user)

		response = self.client.get("/")

		self.assertEqual(response.status_code, 200)
		self.assertContains(response, "legacy-secret-answer")
		self.assertNotContains(response, "Filter Cards")
		self.assertNotContains(response, 'name="category"')
		with self.assertRaises(NoReverseMatch):
			reverse("fiter-categories")


class CardImportTests(TestCase):
	def setUp(self):
		self.user = get_user_model().objects.create_user("importer")
		self.client.force_login(self.user)

	def upload(self, filename, content):
		return self.client.post(
			reverse("card-import"),
			{
				"file": SimpleUploadedFile(filename, content),
			},
		)

	def test_csv_upload_creates_owned_cards_with_any_language_question(self):
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="hello",
			translation="مرحبا",
			vocalized_translation="مَرْحَبًا",
		)
		response = self.upload(
			"cards.csv",
			b"question,answer\nhello,\xd9\x85\xd8\xb1\xd8\xad\xd8\xa8\xd8\xa7\nbye,goodbye\n",
		)

		self.assertRedirects(response, reverse("card-list"))
		self.assertEqual(list(response.wsgi_request._messages)[0].message, "Imported 2 cards.")
		cards = list(Card.objects.filter(user=self.user).order_by("id"))
		self.assertEqual(len(cards), 2)
		self.assertEqual(cards[0].question, "hello")
		self.assertEqual(cards[0].answer, "مرحبا")
		self.assertEqual(cards[0].answer_vocalized, "مَرْحَبًا")
		self.assertEqual(cards[0].box, 1)
		self.assertEqual(cards[1].answer_vocalized, "")
		other_user = get_user_model().objects.create_user("other-deck")
		self.assertEqual(Card.objects.filter(user=other_user).count(), 0)

	def test_csv_upload_leaves_arabic_vocalization_blank_without_dictionary_entry(self):
		answer = "هي تحب ان تذهب"
		response = self.upload(
			"arabic-cards.csv",
			f"question,answer\nphrase,{answer}\n".encode("utf-8"),
		)

		self.assertRedirects(response, reverse("card-list"))
		card = Card.objects.get(user=self.user, question="phrase")
		self.assertEqual(card.answer_vocalized, "")

	def test_xlsx_upload_creates_cards_with_unrestricted_question_text(self):
		from openpyxl import Workbook

		workbook = Workbook()
		worksheet = workbook.active
		worksheet.append(["question", "answer"])
		worksheet.append(["maison", "house"])
		content = BytesIO()
		workbook.save(content)

		response = self.upload("cards.xlsx", content.getvalue())

		self.assertRedirects(response, reverse("card-list"))
		card = Card.objects.get(user=self.user)
		self.assertEqual(card.question, "maison")
		self.assertEqual(card.answer, "house")

	def test_exported_xlsx_round_trips_question_and_answer(self):
		Card.objects.create(
			user=self.user,
			answer="مرحبا",
			question="hello",
		)
		Card.objects.create(
			user=self.user,
			answer="Haus",
			question="maison",
		)
		other_user = get_user_model().objects.create_user("export-recipient")
		Card.objects.create(
			user=other_user,
			answer="private",
			question="private question",
		)

		export_response = self.client.get(reverse("card-export"))

		self.assertEqual(export_response.status_code, 200)
		self.assertEqual(
			export_response["Content-Disposition"],
			'attachment; filename="flashcards.xlsx"',
		)
		self.assertEqual(
			export_response["Content-Type"],
			"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
		)
		from openpyxl import load_workbook
		exported_workbook = load_workbook(
			BytesIO(export_response.content),
			read_only=True,
			data_only=True,
		)
		exported_rows = list(
			exported_workbook.active.iter_rows(values_only=True)
		)
		exported_workbook.close()
		self.assertEqual(
			exported_rows,
			[
				("question", "answer"),
				("hello", "مرحبا"),
				("maison", "Haus"),
			],
		)
		self.assertNotIn(
			("private question", "private"),
			exported_rows,
		)

		recipient_client = self.client_class()
		recipient_client.force_login(other_user)
		import_response = recipient_client.post(
			reverse("card-import"),
			{
				"file": SimpleUploadedFile(
					"flashcards.xlsx",
					export_response.content,
					content_type=(
						"application/vnd.openxmlformats-officedocument."
						"spreadsheetml.sheet"
					),
				),
			},
		)

		self.assertRedirects(import_response, reverse("card-list"))
		imported_cards = Card.objects.filter(user=other_user).order_by("id")
		self.assertEqual(imported_cards.count(), 3)
		self.assertEqual(
			list(imported_cards.values_list("question", "answer")),
			[
				("private question", "private"),
				("hello", "مرحبا"),
				("maison", "Haus"),
			],
		)

	def test_create_and_import_pages_show_all_box_counts(self):
		Card.objects.create(
			user=self.user,
			answer="مرحبا",
			question="hello",
			box=1,
		)

		for url in (reverse("card-create"), reverse("card-import")):
			with self.subTest(url=url):
				response = self.client.get(url)
				self.assertContains(response, "My cards")
				self.assertContains(response, "<mark>1</mark>")
				self.assertContains(response, "<mark>0</mark>")

	def test_duplicate_rows_are_skipped_and_reported_after_import(self):
		response = self.upload(
			"cards.csv",
			b"question,answer\nhello,world\nhello,another answer\nbye,goodbye\n",
		)

		self.assertRedirects(response, reverse("card-list"))
		self.assertEqual(Card.objects.filter(user=self.user).count(), 2)
		message = list(response.wsgi_request._messages)[0]
		self.assertEqual(
			message.message,
			"Imported 2 cards. Skipped 1 duplicate card: hello.",
		)
		self.assertTrue(message.tags.startswith("warning"))

	def test_invalid_csv_imports_nothing(self):
		response = self.upload("bad.csv", b"front,back\nhello,world\n")
		self.assertEqual(response.status_code, 200)
		self.assertContains(
			response,
			"question and answer",
		)
		self.assertEqual(Card.objects.filter(user=self.user).count(), 0)

		response = self.upload(
			"legacy-language.csv",
			b"question,answer,language\nhello,world,en\n",
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, "question and answer")
		self.assertEqual(Card.objects.filter(user=self.user).count(), 0)

		response = self.upload(
			"invalid-row.csv",
			b"question,answer\nvalid question,valid answer\n,missing question\n",
		)
		self.assertEqual(response.status_code, 200)
		self.assertContains(response, "Row 3")
		self.assertEqual(Card.objects.filter(user=self.user).count(), 0)

	def test_upload_skips_existing_duplicate_but_imports_other_rows(self):
		Card.objects.create(
			user=self.user,
			answer="world",
			question="hello",
		)
		other_user = get_user_model().objects.create_user("other-importer")
		Card.objects.create(
			user=other_user,
			answer="world",
			question="other question",
		)

		response = self.upload(
			"cards.csv",
			b"question,answer\nhello,different translation\nnew question,new answer\n",
		)

		self.assertRedirects(response, reverse("card-list"))
		self.assertEqual(Card.objects.filter(user=self.user).count(), 2)
		self.assertEqual(Card.objects.filter(user=other_user).count(), 1)
		message = list(response.wsgi_request._messages)[0]
		self.assertEqual(
			message.message,
			"Imported 1 card. Skipped 1 duplicate card: hello.",
		)

	def test_upload_requires_login(self):
		self.client.logout()
		response = self.upload("cards.csv", b"question,answer\nhello,world\n")

		self.assertEqual(response.status_code, 302)
		self.assertIn(reverse("login"), response.url)


class ArabicVocalizationTests(TestCase):
	def test_json_vocalization_is_paired_with_arabic_text_on_both_faces(self):
		user = get_user_model().objects.create_user("arbitrary-arabic-owner")
		self.client.force_login(user)
		question = "لماذا لا تأتي إلى البيت"
		answer = "هي تحب ان تذهب"
		question_vocalized = "لِمَاذَا لَا تَأْتِي إِلَى الْبَيْتِ"
		answer_vocalized = "هِيَ تُحِبُّ أَنْ تَذْهَبَ"
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Why do you not come home",
			translation=question,
			vocalized_translation=question_vocalized,
		)
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="She likes to go",
			translation=answer,
			vocalized_translation=answer_vocalized,
		)
		card = Card.objects.create(user=user, question=question, answer=answer)
		card.date_moved = timezone.now() - timedelta(days=2)
		card.save(update_fields=["date_moved"])

		for flip in (None, "1"):
			response = self.client.get(
				reverse("todo-list"),
				{"flip": flip} if flip else {},
			)

			self.assertEqual(
				arabic_display_texts(response, f"question-{card.id}"),
				[f"{question} {question_vocalized}"],
			)
			self.assertEqual(
				arabic_display_texts(response, f"answer-{card.id}"),
				[f"{answer} {answer_vocalized}"],
			)

	def test_manually_entered_vocalization_is_saved_with_arabic_answer(self):
		answer = "هي تحب ان تذهب"
		vocalized_answer = "هِيَ تُحِبُّ أَنْ تَذْهَبَ"
		form = CardForm(data={
			"question": "Arabic phrase",
			"answer": answer,
			"answer_vocalized": vocalized_answer,
			"box": 1,
		})

		self.assertTrue(form.is_valid())
		card = form.save()

		self.assertEqual(card.answer_vocalized, vocalized_answer)

	def test_repeat_card_shows_vocalized_question_before_revealing_answer(self):
		user = get_user_model().objects.create_user("arabic-question-owner")
		self.client.force_login(user)
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Sick",
			translation="مريض",
			vocalized_translation="مَرِيضٌ",
		)
		card = Card.objects.create(
			user=user,
			question="مريض",
			answer="Sick",
		)
		card.date_moved = timezone.now() - timedelta(days=2)
		card.save(update_fields=["date_moved"])

		response = self.client.get(reverse("todo-list"))

		self.assertContains(response, "مريض")
		self.assertContains(response, "مَرِيضٌ")
		self.assertContains(response, "Reveal Answer")
		self.assertContains(response, 'name="flip"')
		self.assertNotContains(response, 'name="flip" value="1" checked')
		self.assertContains(response, 'class="flip-face flip-front"')
		self.assertContains(response, 'class="flip-face flip-back"')
		self.assertContains(response, 'class="flip-toggle flip-forward"')
		self.assertContains(response, "Reveal Answer")
		self.assertContains(response, 'class="flip-toggle flip-return"')
		self.assertContains(response, "Show Question")
		self.assertContains(response, 'aria-label="Reveal answer"')
		front = response.content.decode().split(
			'class="flip-face flip-front"',
			1,
		)[1].split('class="flip-face flip-back"', 1)[0]
		back = response.content.decode().split(
			'class="flip-face flip-back"',
			1,
		)[1].split("</section>", 1)[0]
		self.assertIn('class="card-side-label">QUESTION', front)
		self.assertIn('class="card-side-label">ANSWER', back)
		self.assertIn("مريض", front)
		self.assertIn("مَرِيضٌ", front)
		self.assertIn("Sick", back)
		self.assertEqual(card.answer_vocalized, "")

	def test_flipped_repeat_card_shows_answer_first_and_question_after_reveal(self):
		user = get_user_model().objects.create_user("flipped-arabic-question-owner")
		self.client.force_login(user)
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Sick",
			translation="مريض",
			vocalized_translation="مَرِيضٌ",
		)
		card = Card.objects.create(
			user=user,
			question="مريض",
			answer="Sick",
		)
		card.date_moved = timezone.now() - timedelta(days=2)
		card.save(update_fields=["date_moved"])

		response = self.client.get(reverse("todo-list"), {"flip": "1"})
		rendered = response.content.decode()

		self.assertContains(response, 'name="flip"')
		self.assertContains(response, "checked")
		self.assertContains(response, "Reveal Answer")
		self.assertNotContains(response, "Reveal Question")
		flipped_front = rendered.split(
			'class="flip-face flip-front"',
			1,
		)[1].split('class="flip-face flip-back"', 1)[0]
		flipped_back = rendered.split(
			'class="flip-face flip-back"',
			1,
		)[1].split("</section>", 1)[0]
		self.assertIn('class="card-side-label">QUESTION', flipped_front)
		self.assertIn('class="card-side-label">ANSWER', flipped_back)
		self.assertIn("Sick", flipped_front)
		self.assertIn("مريض", flipped_back)
		self.assertIn("مَرِيضٌ", flipped_back)

	def test_arabic_answer_vocalization_is_shown_on_both_repeat_sides(self):
		user = get_user_model().objects.create_user("arabic-answer-owner")
		self.client.force_login(user)
		DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Sick",
			translation="مريض",
			vocalized_translation="مَرِيضٌ",
		)
		card = Card.objects.create(
			user=user,
			question="Sick",
			answer="مريض",
		)
		card.date_moved = timezone.now() - timedelta(days=2)
		card.save(update_fields=["date_moved"])

		default_response = self.client.get(reverse("todo-list"))
		default_rendered = default_response.content.decode()
		self.assertContains(default_response, "مَرِيضٌ")
		default_back = default_rendered.split(
			'class="flip-face flip-back"',
			1,
		)[1].split("</section>", 1)[0]
		self.assertIn('class="card-side-label">ANSWER', default_back)
		self.assertIn("مريض", default_back)
		self.assertIn("مَرِيضٌ", default_back)

		flipped_response = self.client.get(reverse("todo-list"), {"flip": "1"})
		flipped_rendered = flipped_response.content.decode()
		self.assertContains(flipped_response, "مَرِيضٌ")
		flipped_front = flipped_rendered.split(
			'class="flip-face flip-front"',
			1,
		)[1].split('class="flip-face flip-back"', 1)[0]
		self.assertIn('class="card-side-label">QUESTION', flipped_front)
		self.assertIn("مريض", flipped_front)
		self.assertIn("مَرِيضٌ", flipped_front)

	def test_non_arabic_answers_are_not_vocalized(self):
		form = CardForm(data={
			"question": "Haus",
			"answer": "maison",
			"box": 1,
		})

		self.assertTrue(form.is_valid())
		card = form.save()

		self.assertEqual(card.answer_vocalized, "")

	def test_card_result_displays_vocalized_answer_right_to_left(self):
		entry = DictionaryEntry.objects.create(
			language_pair="en-ar",
			headword="Apple",
			translation="التفاحة",
			vocalized_translation="التُّفَّاحَةُ",
		)
		form = CardForm(data={
			"question": "apple",
			"answer": "التفاحة",
			"dictionary_entry_id": entry.id,
			"box": 1,
		})
		self.assertTrue(form.is_valid())
		card = form.save()
		card.answer_dictionary_vocalized = entry.vocalized_translation

		rendered = render_to_string("cards/card.html", {
			"card": card,
			"check_card": False,
		})

		self.assertIn("التفاحة", rendered)
		self.assertIn("التُّفَّاحَةُ", rendered)
		self.assertIn('lang="ar" dir="rtl"', rendered)

	def test_import_vocalized_dictionary_stores_both_translations(self):
		payload = {
			"metadata": {
				"format": "freedict-vocalized-json-v1",
				"language_pair": "en-ar",
				"entry_count": 1,
			},
			"entries": [{
				"headword": "Apple",
				"translation": "التفاحة",
				"vocalized_translation": "التُّفَّاحَةُ",
			}],
		}
		with TemporaryDirectory() as directory:
			json_path = f"{directory}/dictionary.json"
			with open(json_path, "w", encoding="utf-8") as dictionary_file:
				json.dump(payload, dictionary_file, ensure_ascii=False)
			call_command("import_vocalized_dictionary", json_path, stdout=StringIO())

		entry = DictionaryEntry.objects.get(language_pair="en-ar", headword="Apple")
		self.assertEqual(entry.translation, "التفاحة")
		self.assertEqual(entry.vocalized_translation, "التُّفَّاحَةُ")


class ConsolidateQuestionMigrationTests(TransactionTestCase):
	latest_migration = [(
		"cards",
		"0014_remove_card_category_and_account_category_settings",
	)]

	def test_migration_handles_collisions_and_restores_questions_on_rollback(self):
		before_migration = [("cards", "0012_remove_card_uniq_card_user_answer_de_and_more")]
		after_migration = [("cards", "0013_consolidate_question_field")]

		try:
			MigrationExecutor(connection).migrate(before_migration)
			old_apps = MigrationExecutor(connection).loader.project_state(
				before_migration
			).apps
			User = old_apps.get_model("auth", "User")
			LegacyCard = old_apps.get_model("cards", "Card")
			user = User.objects.create(username="question-migration-owner")
			LegacyCard.objects.create(
				user_id=user.pk,
				answer="translation",
				question_de="same question",
				question_en="",
				question_pl="",
				question_it="",
				question_pt="",
				question_es="",
				language="de",
			)
			LegacyCard.objects.create(
				user_id=user.pk,
				answer="translation",
				question_de="same question (2)",
				question_en="",
				question_pl="",
				question_it="",
				question_pt="",
				question_es="",
				language="de",
			)
			LegacyCard.objects.create(
				user_id=user.pk,
				answer="other translation",
				question_de="",
				question_en="same question",
				question_pl="",
				question_it="",
				question_pt="",
				question_es="",
				language="en",
			)

			MigrationExecutor(connection).migrate(after_migration)
			current_apps = MigrationExecutor(connection).loader.project_state(
				after_migration
			).apps
			CurrentCard = current_apps.get_model("cards", "Card")
			self.assertEqual(
				list(
					CurrentCard.objects.filter(user_id=user.pk)
					.order_by("pk")
					.values_list(
						"question",
						flat=True,
					)
				),
				["same question", "same question (2)", "same question (3)"],
			)
			with self.assertRaises(IntegrityError), transaction.atomic():
				CurrentCard.objects.create(
					user_id=user.pk,
					question="same question",
					answer="a different answer",
				)

			MigrationExecutor(connection).migrate(before_migration)
			restored_apps = MigrationExecutor(connection).loader.project_state(
				before_migration
			).apps
			RestoredCard = restored_apps.get_model("cards", "Card")
			restored_rows = list(
				RestoredCard.objects.filter(user_id=user.pk).values_list(
					"question_de",
					"question_en",
					"question_pl",
					"question_it",
					"question_pt",
					"question_es",
					"language",
				)
			)
			self.assertEqual(len(restored_rows), 3)
			self.assertTrue(
				all(any(row[:6]) for row in restored_rows)
			)
			self.assertEqual(
				{row[0] for row in restored_rows},
				{
					"same question",
					"same question (2)",
					"same question (3)",
				},
			)
			self.assertEqual(
				{row[6] for row in restored_rows},
				{"de"},
			)
		finally:
			MigrationExecutor(connection).migrate(self.latest_migration)

	def test_migration_removes_category_field_and_account_settings(self):
		before_migration = [("cards", "0013_consolidate_question_field")]

		try:
			MigrationExecutor(connection).migrate(before_migration)
			old_apps = MigrationExecutor(connection).loader.project_state(
				before_migration
			).apps
			User = old_apps.get_model("auth", "User")
			LegacyCard = old_apps.get_model("cards", "Card")
			CategorySettings = old_apps.get_model("cards", "AccountCategorySettings")
			user = User.objects.create(username="category-migration-owner")
			CategorySettings.objects.create(user_id=user.pk, arabic_only=True)
			LegacyCard.objects.create(
				user_id=user.pk,
				answer="translation",
				question="question",
				category="Arabic",
			)

			MigrationExecutor(connection).migrate(self.latest_migration)
			current_apps = MigrationExecutor(connection).loader.project_state(
				self.latest_migration
			).apps
			CurrentCard = current_apps.get_model("cards", "Card")
			self.assertNotIn("category", [field.name for field in CurrentCard._meta.fields])
			self.assertNotIn(
				"accountcategorysettings",
				current_apps.all_models["cards"],
			)
			with connection.cursor() as cursor:
				card_columns = {
					column.name
					for column in connection.introspection.get_table_description(
						cursor,
						CurrentCard._meta.db_table,
					)
				}
				table_names = connection.introspection.table_names(cursor)
			self.assertNotIn("category", card_columns)
			self.assertNotIn("cards_accountcategorysettings", table_names)
		finally:
			MigrationExecutor(connection).migrate(self.latest_migration)
